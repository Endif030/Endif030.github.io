from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# 创建演示文稿 - 恋爱女友主题
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 恋爱女友主题数据
gf_jokes = [
    {
        "id": "2.1",
        "title": "免责声明",
        "text": "我今晚要讲关于我女朋友的事，但我想小心一点，因为很多男性喜剧演员会注意到他们女朋友或妻子的某些特质，然后在舞台上表现得好像这个特质适用于整个性别。所以我想明确一点，接下来的笑话是关于我女朋友的。我不想用她的真名，所以我们叫她 women 吧。",
        "relation": "转入女友话题"
    },
    {
        "id": "2.2",
        "title": "女友是犹太人",
        "text": "Women are Jewish（女人是犹太人），我的也不例外。事实上，我女朋友在 Hasidic（正统犹太教）社区长大。你认识 Hasids 吗？纽约有很多，戴着帽子，留着鬓角。他们总是在火车上看 Torah。每次看到正统犹太教徒看 Torah，我就走过去说：你怎么还没读完？天啊。请在买服装之前看完。你看起来疯了。",
        "relation": "承接2.1：展开女友背景"
    },
    {
        "id": "2.3",
        "title": "色盲言论",
        "text": "没关系，我是犹太人。我不觉得我需要这么说，但是你知道，现在有些人会说：我看不见鼻子（I don't see noses）。",
        "relation": "承接2.2：犹太人身份"
    },
    {
        "id": "2.4",
        "title": "文化犹太人定义",
        "text": "我不是在很有宗教氛围的环境中长大的，但我认为自己是一个文化上的犹太人（culturally Jewish），这意味着我有普通犹太教的所有肠胃问题、压力和焦虑，但没有上帝甜蜜的安慰（without the sweet comfort of God）。",
        "relation": "承接2.3：犹太人身份展开"
    },
    {
        "id": "2.5",
        "title": "逾越节后的疑惑",
        "text": "我庆祝大多数主要的犹太节日。逾越节是我最喜欢的，因为你可以在餐桌上读故事。那是犹太人被法老奴役，然后摩西让瘟疫发生，然后我们逃脱并穿越沙漠40年的故事。我喜欢它，因为它让犹太人看起来很酷，但这也让我有点困惑，因为我们是怎么从穿越沙漠变成现在被称为 generally frail（通常体弱）的呢？",
        "relation": "承接2.4：犹太文化展开"
    },
    {
        "id": "2.6",
        "title": "迈阿密度假发现",
        "text": "直到有一天，我和我的犹太女友去迈阿密度假。我们一到海滩，她就说：这里像沙漠一样。我说：天啊，如果犹太人一直在 exaggerating（夸张）这一切呢？我们在路上堵住了。我女朋友说：let my people go（让我的人民走）。海滩有点拥挤。她说：这里有六百万（6 million）人。",
        "relation": "承接2.5：犹太女友娇气表现"
    },
    {
        "id": "2.7",
        "title": "基督教喜剧",
        "text": "是的，你懂了。你想要可爱的笑话？去看基督教喜剧演员（Christian comedian）。",
        "relation": "承接2.6：自嘲过渡"
    },
    {
        "id": "2.8",
        "title": "宗教社区和友谊",
        "text": "我确实羡慕宗教人士的社区感。直到我年纪大了才意识到，但一到三十多岁，我就意识到和朋友找时间相聚有多难。有时我周五晚上有空，我会联系朋友，但他们很忙，反之亦然。那就是 Shabbat dinner（安息日晚餐）的作用。犹太人决定每周五需要一个固定的约会，这样他们就能在一起。",
        "relation": "承接2.7：宗教社区观察"
    },
    {
        "id": "2.9",
        "title": "维持友谊的秘诀",
        "text": "看看穆斯林，他们每天应该祈祷五次，但很多时候是一起祈祷。尽管我不相信这些，但我不得不承认，维持友谊的关键似乎是 having hating gay people（一起恨同性恋）。",
        "relation": "承接2.8：友谊话题延伸"
    },
    {
        "id": "2.10",
        "title": "女友是第一任正式关系",
        "text": "有时我觉得我有优势，因为我是我女朋友的第一段正式关系。事实上，她对我说话的方式只在第一段正式关系中才会出现。有时晚上，她会转向我说：you're my person（你是我的那个人）。我说：you're last but not least（你最后但并非最不重要）。",
        "relation": "转入女友关系话题"
    },
    {
        "id": "2.11",
        "title": "前女友分手场景",
        "text": "我在大学遇到一个女人，我们约会了五年。五年后，她在中午繁忙的星巴克和我分手，因为她不想我大吵大闹。我说：那你不应该带观众来（brought an audience）。",
        "relation": "承接2.10：前任话题"
    },
    {
        "id": "2.12",
        "title": "纹身Larry",
        "text": "我纹了她名字的第一个字母。没成功，她嫁给了别人。我不得不告诉每个约会对象关于我死去的朋友 Larry 的事。\\n【梗：L 纹身伪装成纪念死去的朋友 Larry】",
        "relation": "承接2.11：纹身故事"
    },
    {
        "id": "2.13",
        "title": "用专场换纹身",
        "text": "我决定妥协。我说：让我想办法拍一个喜剧专场，在胶片上讲所有关于纹身的笑话。我向她保证，如果这赢不回 Laura，我第二天就去掉。\\n【callback：专场名 Thief of Joy 的来源】",
        "relation": "承接2.12：纹身冲突升级"
    },
    {
        "id": "2.14",
        "title": "纹身的性影响",
        "text": "我女朋友不给我口交，因为她说看到纹身会想起我前任。我希望我纹在背上，这样她就不会一直试图 peg（肛交）我了。",
        "relation": "承接2.13：纹身后果"
    },
    {
        "id": "2.15",
        "title": "女友的夜惊症",
        "text": "睡眠很困难，因为我女朋友有夜惊症（night terrors）。半夜，我女朋友突然会说：有人闯进来了。我说：在哪里？她说完就睡着了。但我有慢性焦虑（chronic anxiety），所以我整晚都醒着，成为我们两个精神疾病的唯一受害者。",
        "relation": "转入女友夜惊话题"
    },
    {
        "id": "2.16",
        "title": "夜惊的强奸犯类比",
        "text": "她完全不记得夜惊。所以早上她翻过身问：你睡得好吗？就像强奸犯问 was it good for you（你爽吗）？",
        "relation": "承接2.15：夜惊后果"
    },
    {
        "id": "2.17",
        "title": "夜惊导致错过真事",
        "text": "有一次半夜，我女朋友从床上爬起来开始收拾行李。我看着手机，凌晨四点。我说：宝贝，你在做夜惊。她说：不，我刚收到短信说我爸在医院。我得马上去机场。我拿起手机说：哪个机场？她说... son of a bitch。",
        "relation": "承接2.16：夜惊故事"
    },
    {
        "id": "2.18",
        "title": "要不要孩子的反问",
        "text": "我不知道是否想要孩子。天啊，人们一直问我。我初次见到人，如果和女朋友一起，他们会在一分钟内问：你们想过要孩子吗？我说：这不是闲聊。你打算怎么处理你的遗体（doing with your remains）？",
        "relation": "转入孩子话题"
    },
    {
        "id": "2.19",
        "title": "孩子=袋鼠",
        "text": "我住在纽约，从事娱乐业，我没有有孩子的朋友。所以问我要不要孩子就像问我要不要袋鼠（kangaroo）。是的，我熟悉，但我不知道怎么打。",
        "relation": "承接2.18：孩子话题展开"
    },
    {
        "id": "2.20",
        "title": "冻卵和狼人",
        "text": "我女朋友去年冷冻了卵子，这是我的主意。医生展示了她需要注射的一长串针剂。他说：最坏的情况，就像 PMS 的10倍。我说：我不知道那也是计量单位。他说：因为每天她都需要注射雌激素，这就像和狼人约会并注射月亮（injecting the moon）。",
        "relation": "承接2.19：冻卵话题"
    },
    {
        "id": "2.21",
        "title": "冻卵和科学奇迹",
        "text": "我不知道一个女人要有事业需要多少科学奇迹（scientific miracles）。",
        "relation": "承接2.20：冻卵话题总结"
    },
    {
        "id": "2.22",
        "title": "阿拉巴马胚胎裁决",
        "text": "我们得到了卵子。我很欣慰，直到我读到阿拉巴马那个该死的法官裁定胚胎现在被认为是人。我们国家的发展方向，再有两个裁决我女朋友就要成为19个孩子的妈了（mother of 19）。",
        "relation": "承接2.21：政治话题"
    },
    {
        "id": "2.23",
        "title": "镜子上的床",
        "text": "我真的想在床上方安装一面镜子，但这只是为了证明我女朋友占了超过一半（taking more than half）。",
        "relation": "转入关系观察"
    }
]

# 颜色定义
BG_COLOR = RGBColor(107, 155, 209)  # 蓝色
BORDER_COLOR = RGBColor(26, 82, 118)
TITLE_COLOR = RGBColor(21, 67, 96)
TEXT_COLOR = RGBColor(26, 82, 118)
RELATION_COLOR = RGBColor(41, 128, 185)

# 每页放2个笑点
jokes_per_slide = 2

for slide_idx in range(0, len(gf_jokes), jokes_per_slide):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    if slide_idx == 0:
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.6))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = "【恋爱女友】主题笑点详解 - Tom Segura Thief of Joy"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TITLE_COLOR
        p.alignment = PP_ALIGN.CENTER
        start_y = 1.0
    else:
        start_y = 0.4
    
    for i, joke in enumerate(gf_jokes[slide_idx:slide_idx + jokes_per_slide]):
        y_pos = start_y + i * 3.4
        
        shape = slide.shapes.add_shape(1, Inches(0.5), Inches(y_pos), Inches(12.333), Inches(3.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = BG_COLOR
        shape.line.color.rgb = BORDER_COLOR
        shape.line.width = Pt(3)
        
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_pos + 0.15), Inches(11.9), Inches(0.4))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"{joke['id']} {joke['title']}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TITLE_COLOR
        
        content_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_pos + 0.6), Inches(11.9), Inches(2.0))
        tf = content_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = joke['text']
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_COLOR
        p.line_spacing = 1.3
        
        relation_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_pos + 2.7), Inches(11.9), Inches(0.35))
        tf = relation_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"→ {joke['relation']}"
        p.font.size = Pt(10)
        p.font.color.rgb = RELATION_COLOR
        p.font.italic = True
    
    if slide_idx + jokes_per_slide < len(gf_jokes):
        page_box = slide.shapes.add_textbox(Inches(6), Inches(7.1), Inches(1.333), Inches(0.3))
        tf = page_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"第 {slide_idx//jokes_per_slide + 1} 页 / 共 12 页"
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor(127, 140, 141)
        p.alignment = PP_ALIGN.CENTER

ppt_path = '/root/.openclaw/workspace/transcripts/Thief_of_Joy_恋爱女友.pptx'
prs.save(ppt_path)
print(f"✅ 恋爱女友主题PPT已生成: {ppt_path}")
print(f"共 {len(prs.slides)} 页")
