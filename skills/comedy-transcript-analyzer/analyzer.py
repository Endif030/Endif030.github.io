# Comedy Transcript Analyzer
# 脱口秀逐字稿拆解分析工具

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

class ComedyTranscriptAnalyzer:
    """脱口秀逐字稿分析器"""
    
    # 主题配色方案
    THEME_COLORS = {
        "原生家庭": RGBColor(126, 181, 166),      # 青绿
        "恋爱关系": RGBColor(107, 155, 209),      # 蓝色
        "表演生涯": RGBColor(144, 200, 172),      # 薄荷绿
        "性身体": RGBColor(244, 208, 63),         # 金黄
        "生死心理": RGBColor(241, 148, 138),      # 珊瑚
        "社会观察": RGBColor(174, 214, 241),      # 浅蓝
        "默认": RGBColor(189, 195, 199)
    }
    
    def __init__(self, output_dir="./output"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    def create_ppt(self, title, jokes_data, filename):
        """
        创建PPT文件
        
        Args:
            title: PPT标题
            jokes_data: 笑点数据列表，每个元素为字典 {id, title, text, relation, theme}
            filename: 输出文件名
        """
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        jokes_per_slide = 2
        
        for slide_idx in range(0, len(jokes_data), jokes_per_slide):
            blank_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_layout)
            
            if slide_idx == 0:
                title_box = slide.shapes.add_textbox(
                    Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.6)
                )
                tf = title_box.text_frame
                p = tf.paragraphs[0]
                p.text = title
                p.font.size = Pt(28)
                p.font.bold = True
                p.font.color.rgb = RGBColor(21, 67, 96)
                p.alignment = PP_ALIGN.CENTER
                start_y = 1.0
            else:
                start_y = 0.4
            
            for i, joke in enumerate(jokes_data[slide_idx:slide_idx + jokes_per_slide]):
                y_pos = start_y + i * 3.4
                theme = joke.get("theme", "默认")
                bg_color = self.THEME_COLORS.get(theme, self.THEME_COLORS["默认"])
                
                # 添加形状
                shape = slide.shapes.add_shape(
                    1, Inches(0.5), Inches(y_pos), Inches(12.333), Inches(3.1)
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = bg_color
                shape.line.color.rgb = RGBColor(26, 82, 118)
                shape.line.width = Pt(3)
                
                # 标题
                title_box = slide.shapes.add_textbox(
                    Inches(0.7), Inches(y_pos + 0.15), Inches(11.9), Inches(0.4)
                )
                tf = title_box.text_frame
                p = tf.paragraphs[0]
                p.text = f"{joke['id']} {joke['title']}"
                p.font.size = Pt(16)
                p.font.bold = True
                p.font.color.rgb = RGBColor(21, 67, 96)
                
                # 内容
                content_box = slide.shapes.add_textbox(
                    Inches(0.7), Inches(y_pos + 0.6), Inches(11.9), Inches(2.0)
                )
                tf = content_box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = joke['text']
                p.font.size = Pt(12)
                p.font.color.rgb = RGBColor(26, 82, 118)
                p.line_spacing = 1.3
                
                # 关系说明
                relation_box = slide.shapes.add_textbox(
                    Inches(0.7), Inches(y_pos + 2.7), Inches(11.9), Inches(0.35)
                )
                tf = relation_box.text_frame
                p = tf.paragraphs[0]
                p.text = f"→ {joke.get('relation', '')}"
                p.font.size = Pt(10)
                p.font.color.rgb = RGBColor(41, 128, 185)
                p.font.italic = True
            
            # 页码
            if slide_idx + jokes_per_slide < len(jokes_data):
                page_box = slide.shapes.add_textbox(
                    Inches(6), Inches(7.1), Inches(1.333), Inches(0.3)
                )
                tf = page_box.text_frame
                p = tf.paragraphs[0]
                p.text = f"第 {slide_idx//jokes_per_slide + 1} 页"
                p.font.size = Pt(9)
                p.font.color.rgb = RGBColor(127, 140, 141)
                p.alignment = PP_ALIGN.CENTER
        
        output_path = os.path.join(self.output_dir, filename)
        prs.save(output_path)
        return output_path

# 示例用法
if __name__ == "__main__":
    analyzer = ComedyTranscriptAnalyzer()
    
    # 示例数据
    sample_jokes = [
        {
            "id": "1.1",
            "title": "开场钩子",
            "text": "我来自一个 messed up 的家庭...你们需要专业人士（needed a professional）。",
            "relation": "开场引入",
            "theme": "原生家庭"
        }
    ]
    
    analyzer.create_ppt(
        "【示例】脱口秀笑点分析",
        sample_jokes,
        "示例分析.pptx"
    )
